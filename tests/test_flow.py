"""End-to-end lifecycle through the web UI using the offline stub and demo catalogue."""
import io
import re

from openpyxl import load_workbook
from pptx import Presentation

PAGES = ["/", "/backlog", "/portfolio", "/framework", "/rounds", "/usecases/new", "/api/usecases", "/healthz"]


def test_pages_render(client):
    for url in PAGES:
        assert client.get(url).status_code == 200, url


def test_seeded_worked_example(client):
    data = client.get("/api/usecases").json()
    top = [d["title"] for d in data if d["computed_rank"]][:3]
    assert top == ["Teaching demand forecast", "At-risk student alerts", "Budget variance commentary (GenAI)"]


def _ref_from(resp):
    return re.search(r"/usecases/(UC-\d+)", resp.headers["location"]).group(1)


def test_full_lifecycle(client):
    r = client.post("/usecases/new", data={
        "role": "timetabling officer", "need": "a view of room utilisation against enrolments",
        "outcome": "I can release under-used rooms before the semester",
        "answer_time_lost": "Chasing enrolment numbers from unit chairs every week, manual spreadsheets",
        "answer_six_months": "Rooms matched to enrolments a fortnight before teaching starts",
        "domain": "Teaching and Learning", "then": "view",
    }, follow_redirects=False)
    assert r.status_code == 303
    ref = _ref_from(r)

    # story drafting (HTMX partial) and apply
    r = client.post(f"/usecases/{ref}/ai/story")
    assert r.status_code == 200 and "Suggested story" in r.text
    r = client.post(f"/usecases/{ref}/story/apply", data={"title": "Room utilisation view", "role": "timetabling officer",
                    "need": "a view of room utilisation against enrolments", "outcome": "I can release rooms", "gap_statement": "Gap text"},
                    follow_redirects=False)
    assert r.status_code == 303

    # qualification requires notes when checks are missing
    r = client.post(f"/usecases/{ref}/qualify", data={"outcome": "qualified", "check_named_role": "on"}, follow_redirects=False)
    assert "err=" in r.headers["location"]
    checks = {f"check_{k}": "on" for k in ["named_role", "real_problem", "gap_clear", "lead_identified", "data_solvable"]}
    r = client.post(f"/usecases/{ref}/qualify", data={"outcome": "qualified", **checks}, follow_redirects=False)
    assert "tab=size" in r.headers["location"]

    # catalogue-informed sizing suggestion, then accept a cell
    client.post(f"/usecases/{ref}/ai/size")
    uc = next(d for d in client.get("/api/usecases").json() if d["ref"] == ref)
    page = client.get(f"/usecases/{ref}?tab=size").text
    assert "Data needs and catalogue evidence" in page and "fact_unit_enrolment" in page
    r = client.post(f"/usecases/{ref}/size", data={"cell": "M|M", "rationale": "Two sources"}, follow_redirects=False)
    assert r.status_code == 303

    # scoring: suggestion then human override
    client.post(f"/usecases/{ref}/ai/score")
    r = client.post(f"/usecases/{ref}/score", data={"level_desirability": "H", "level_feasibility": "M", "level_viability": "M"}, follow_redirects=False)
    assert "Priority" in r.headers["location"] or "msg=" in r.headers["location"]
    uc = next(d for d in client.get("/api/usecases").json() if d["ref"] == ref)
    assert uc["sizing"]["size"] == "M" and uc["scores"]["desirability"] == "H" and uc["status"] == "scored"

    # agreed rank moves to prioritised
    client.post(f"/usecases/{ref}/rank", data={"agreed_rank": "2"})
    uc = next(d for d in client.get("/api/usecases").json() if d["ref"] == ref)
    assert uc["agreed_rank"] == 2 and uc["status"] == "prioritised"

    # lean canvas, accelerators
    client.post(f"/usecases/{ref}/ai/canvas")
    uc = next(d for d in client.get("/api/usecases").json() if d["ref"] == ref)
    assert set(uc["canvas"]) >= {"business_problem", "required_data", "estimated_resources"}
    client.post(f"/usecases/{ref}/ai/sql")
    client.post(f"/usecases/{ref}/ai/mockup")
    acc = client.get(f"/usecases/{ref}?tab=accelerators").text
    assert "SELECT" in acc and "iframe" in acc
    mock = client.get(f"/usecases/{ref}/mockup")
    assert "sandbox" in mock.headers["content-security-policy"]

    # audit trail captured each step, including the AI override
    hist = client.get(f"/usecases/{ref}?tab=history").text
    for phrase in ["Captured", "Qualification: qualified", "Sized Medium", "Agreed rank", "Lean Canvas generated"]:
        assert phrase in hist, phrase


def test_park_requires_note(client):
    ref = _ref_from(client.post("/usecases/new", data={"role": "analyst", "need": "x", "then": "view"}, follow_redirects=False))
    r = client.post(f"/usecases/{ref}/qualify", data={"outcome": "parked"}, follow_redirects=False)
    assert "err=" in r.headers["location"]
    client.post(f"/usecases/{ref}/qualify", data={"outcome": "parked", "note": "Needs a named sponsor"})
    assert next(d for d in client.get("/api/usecases").json() if d["ref"] == ref)["status"] == "parked"


def test_notes_import(client):
    notes = """- Workload planners rebuild the tutor allocation spreadsheet every trimester from three exports and enrolment numbers are wrong until census.

- The finance business partner spends two days drafting budget variance commentary each month.

- Something about better dashboards generally."""
    r = client.post("/rounds/1/extract", data={"notes": notes})
    assert r.status_code == 200 and "Candidate use cases" in r.text
    payload = re.search(r'name="payload" value="([^"]*)"', r.text).group(1)
    import html
    r = client.post("/rounds/1/accept", data={"payload": html.unescape(payload), "accept": ["0", "1"]}, follow_redirects=False)
    assert "Added 2 use case" in client.get(r.headers["location"]).text or "Added%202" in r.headers["location"]


def test_exports_open(client):
    wb = load_workbook(io.BytesIO(client.get("/export/backlog.xlsx").content))
    assert wb.active["B1"].value == "Ref" and wb.active.max_row > 3
    prs = Presentation(io.BytesIO(client.get("/export/portfolio.pptx").content))
    assert len(prs.slides) >= 4
    prs = Presentation(io.BytesIO(client.get("/usecases/UC-001/canvas.pptx").content))
    texts = " ".join(sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame)
    assert "Business Problem Statement" in texts and "Estimated Resources" in texts
