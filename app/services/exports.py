"""PowerPoint and Excel exports. Layout follows the Lean Canvas template:
five boxes across the top; benefits/foundations stacked, required data,
assumptions and resources across the bottom."""
from __future__ import annotations

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from ..config import Framework
from ..models import UseCase
from .ranking import RankedItem

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)


def _rgb(hex_colour: str) -> RGBColor:
    return RGBColor.from_string(hex_colour.lstrip("#").upper())


def _text(slide, x, y, w, h, text, size=12, bold=False, colour="#13294b", align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.02)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size, run.font.bold, run.font.color.rgb = Pt(size), bold, _rgb(colour)
    return box


def _body_size(text: str, w: Emu, h: Emu) -> int:
    """Largest point size at which the wrapped lines fit the box (approximate)."""
    width_pt, height_pt = (w / 12700) - 18, (h / 12700) - 6
    lines = [ln for ln in (text or "").splitlines() if ln.strip()] or ["TBC"]
    for size in (11, 10, 9, 8, 7):
        per_line = max(int(width_pt / (size * 0.52)), 8)
        needed = sum(-(-len(ln) // per_line) for ln in lines)
        if needed * size * 1.3 + len(lines) * 2 <= height_pt:
            return size
    return 7


def _box(slide, fw: Framework, x, y, w, h, title: str, body: str):
    c = fw.raw["brand"]["colours"]
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.adjustments[0] = 0.04
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(c["surface"])
    shape.line.color.rgb = _rgb(c["border"])
    shape.line.width = Pt(1.25)
    shape.shadow.inherit = False
    tf = shape.text_frame
    tf.word_wrap = True
    for side in ("margin_left", "margin_right"):
        setattr(tf, side, Inches(0.12))
    tf.margin_top, tf.margin_bottom = Inches(0.08), Inches(0.06)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = title
    run.font.size, run.font.bold, run.font.color.rgb = Pt(12), True, _rgb(c["accent"])
    size = _body_size(body, w, h - Inches(0.4))
    for line in [ln for ln in (body or "").splitlines() if ln.strip()] or ["TBC"]:
        para = tf.add_paragraph()
        para.alignment = PP_ALIGN.LEFT
        para.space_before = Pt(2)
        r = para.add_run()
        text = line.strip()
        r.text = "•\u00a0" + text[2:] if text.startswith("- ") else text
        r.font.size = Pt(size)
        r.font.color.rgb = _rgb("#44546a")


def _canvas_slide(prs: Presentation, fw: Framework, uc: UseCase) -> None:
    c = fw.raw["brand"]["colours"]
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _text(slide, Inches(0.4), Inches(0.25), Inches(9.6), Inches(0.6), f"{fw.raw['lean_canvas']['title']}: {uc.title}", 26, True, c["ink"])
    _text(slide, Inches(0.4), Inches(0.85), Inches(9.6), Inches(0.5), uc.story or fw.raw["lean_canvas"]["subtitle"], 12, False, "#2f5f8a")
    chips = [uc.ref]
    if uc.size:
        chips.append(f"Size {fw.size_by_key[uc.size]['label']}")
    if uc.scores:
        chips.append(" ".join(f"{k[0].upper()}:{v}" for k, v in uc.scores.items()))
    chips.append(fw.status_labels.get(uc.status, uc.status))
    _text(slide, Inches(9.9), Inches(0.3), Inches(3.05), Inches(0.9), "  ·  ".join(chips), 10, True, c["accent"], PP_ALIGN.RIGHT)

    sections = {s["key"]: s for s in fw.canvas_sections}
    content = uc.canvas or {}
    left, top, gap = Inches(0.4), Inches(1.45), Inches(0.1)
    total_w = SLIDE_W - Inches(0.8)
    row1_h, row2_h = Inches(2.6), Inches(2.85)

    def body(key):
        return content.get(key, "")

    def title(key):
        return sections[key]["title"] if key in sections else key

    top_keys = ["business_problem", "solution_description", "systems_data_tech", "capability_requirements", "business_users"]
    weights = [365, 483, 303, 303, 392]
    usable = total_w - gap * (len(top_keys) - 1)
    x = left
    for key, wt in zip(top_keys, weights):
        w = int(usable * wt / sum(weights))
        if key in sections:
            _box(slide, fw, x, top, w, row1_h, title(key), body(key))
        x += w + gap

    y2 = top + row1_h + gap
    bottom = [("stack", 505), ("required_data", 343), ("assumptions_risks", 620), ("estimated_resources", 392)]
    usable = total_w - gap * (len(bottom) - 1)
    x = left
    for key, wt in bottom:
        w = int(usable * wt / sum(b[1] for b in bottom))
        if key == "stack":
            h_top = int(row2_h * 0.48)
            if "estimated_benefits" in sections:
                _box(slide, fw, x, y2, w, h_top, title("estimated_benefits"), body("estimated_benefits"))
            if "data_foundations" in sections:
                _box(slide, fw, x, y2 + h_top + gap, w, row2_h - h_top - gap, title("data_foundations"), body("data_foundations"))
        elif key in sections:
            _box(slide, fw, x, y2, w, row2_h, title(key), body(key))
        x += w + gap

    extra = [s for s in fw.canvas_sections if s["key"] not in set(top_keys) | {"estimated_benefits", "data_foundations", "required_data", "assumptions_risks", "estimated_resources"}]
    if extra:  # client-added sections go on a continuation slide
        s2 = prs.slides.add_slide(prs.slide_layouts[6])
        _text(s2, Inches(0.4), Inches(0.25), Inches(12), Inches(0.6), f"{uc.title}: additional canvas sections", 22, True, c["ink"])
        for i, s in enumerate(extra):
            _box(s2, fw, Inches(0.4) + (i % 3) * Inches(4.2), Inches(1.2) + (i // 3) * Inches(3.0), Inches(4.0), Inches(2.8), s["title"], body(s["key"]))

    _text(slide, Inches(0.4), SLIDE_H - Inches(0.4), Inches(12.5), Inches(0.3),
          f"{fw.raw['brand']['product_name']} · {uc.ref} · exported {date.today():%d %b %Y} · AI-assisted draft, reviewed by the use case lead",
          8, False, "#8a96a8")


def lean_canvas_pptx(fw: Framework, use_cases: list[UseCase]) -> bytes:
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    for uc in use_cases:
        _canvas_slide(prs, fw, uc)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _table(slide, x, y, w, rows: list[list[str]], col_w: list[float], fw: Framework, font=11):
    c = fw.raw["brand"]["colours"]
    shape = slide.shapes.add_table(len(rows), len(rows[0]), x, y, w, Inches(0.4) * len(rows))
    table = shape.table
    for i, cw in enumerate(col_w):
        table.columns[i].width = int(w * cw)
    for r, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.cell(r, ci)
            cell.text = str(val)
            para = cell.text_frame.paragraphs[0]
            para.font.size = Pt(font)
            para.font.bold = r == 0
            para.font.color.rgb = _rgb("#ffffff" if r == 0 else c["ink"])
            cell.fill.solid()
            cell.fill.fore_color.rgb = _rgb(c["ink"] if r == 0 else ("#ffffff" if r % 2 else c["surface"]))
    return table


def portfolio_pptx(fw: Framework, ranked: list[RankedItem], title: str) -> bytes:
    c = fw.raw["brand"]["colours"]
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    labels = fw.level_labels

    s = prs.slides.add_slide(prs.slide_layouts[6])
    _text(s, Inches(0.6), Inches(2.6), Inches(12), Inches(1), title, 36, True, c["ink"])
    _text(s, Inches(0.6), Inches(3.6), Inches(12), Inches(0.6),
          f"Use case portfolio · {date.today():%d %B %Y} · {len(ranked)} use cases", 16, False, "#2f5f8a")

    scored = [r for r in ranked if r.score is not None]
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _text(s, Inches(0.5), Inches(0.3), Inches(12), Inches(0.6), "Scoring on desirability, feasibility and viability", 26, True, c["ink"])
    _text(s, Inches(0.5), Inches(0.9), Inches(12), Inches(0.4),
          "Ranked by weighted DFV, then smaller size, then desirability. Agreed ranks from portfolio review take precedence.", 12, False, "#2f5f8a")
    rows = [["Priority", "Use case", *[lens["label"] for lens in fw.lenses], "Size", "Flags"]]
    for item in scored[:14]:
        uc = item.use_case
        rows.append([str(item.computed_rank), f"{uc.ref} {uc.title}", *[labels.get(uc.scores.get(k), "-") for k in fw.lens_keys],
                     fw.size_by_key.get(uc.size or "", {}).get("label", "-"), "; ".join(item.flags)])
    if len(rows) == 1:
        rows.append(["-", "No scored use cases yet", "", "", "", "", ""])
    _table(s, Inches(0.5), Inches(1.5), Inches(12.3), rows, [0.07, 0.33, 0.1, 0.1, 0.1, 0.08, 0.22], fw, 11)

    # sizing matrix
    s = prs.slides.add_slide(prs.slide_layouts[6])
    sz = fw.raw["sizing"]
    _text(s, Inches(0.5), Inches(0.3), Inches(12), Inches(0.6), "Size and effort read", 26, True, c["ink"])
    _text(s, Inches(0.5), Inches(0.9), Inches(12), Inches(0.4), sz["next_step"], 12, False, "#2f5f8a")
    gx, gy, cw, ch = Inches(2.6), Inches(1.9), Inches(3.0), Inches(1.6)
    for ci, col in enumerate(sz["columns"]["levels"]):
        _text(s, gx + ci * cw, gy - Inches(0.55), cw, Inches(0.5), f"{col['label']}\n{col['description']}", 10, True, c["ink"], PP_ALIGN.CENTER)
    for ri, row in enumerate(sz["rows"]["levels"]):
        _text(s, Inches(0.5), gy + ri * ch + Inches(0.2), Inches(2.0), ch, f"{row['label']}\n{row['description']}", 10, True, c["ink"])
        for ci, col in enumerate(sz["columns"]["levels"]):
            size_key = sz["matrix"][row["key"]][col["key"]]
            size = fw.size_by_key[size_key]
            cell = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, gx + ci * cw + Inches(0.04), gy + ri * ch + Inches(0.04), cw - Inches(0.08), ch - Inches(0.08))
            cell.adjustments[0] = 0.06
            cell.fill.solid()
            cell.fill.fore_color.rgb = _rgb(size["colour"])
            cell.line.fill.background()
            refs = [i.use_case.ref for i in ranked if i.use_case.scope_level == row["key"] and i.use_case.effort_level == col["key"]]
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            r.text = size["label"]
            r.font.size, r.font.bold, r.font.color.rgb = Pt(16), True, _rgb("#ffffff")
            if refs:
                p2 = tf.add_paragraph()
                p2.alignment = PP_ALIGN.CENTER
                r2 = p2.add_run()
                r2.text = ", ".join(refs[:8]) + (" …" if len(refs) > 8 else "")
                r2.font.size, r2.font.color.rgb = Pt(10), _rgb("#ffffff")
    for i, size in enumerate(fw.sizes):
        _text(s, Inches(11.8), gy + i * Inches(0.9), Inches(1.4), Inches(0.9), f"{size['label']}\n{size['weeks']}", 9, False, c["ink"])

    for item in ranked:
        if item.use_case.canvas:
            _canvas_slide(prs, fw, item.use_case)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def backlog_xlsx(fw: Framework, ranked: list[RankedItem]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Backlog"
    headers = ["Rank", "Ref", "Title", "Status", "Stage", "Domain", "Use case lead", "Sponsor", "Story", "Gap statement",
               "Solution type", "Scope", "Effort", "Size", "Indicative weeks", *[lens["label"] for lens in fw.lenses],
               "Priority score", "Agreed rank", "Flags", "Qualification", "Updated"]
    ws.append(headers)
    sol = {s["key"]: s["label"] for s in fw.raw["solution_types"]}
    for item in ranked:
        uc = item.use_case
        stage = fw.stage_for_status(uc.status)
        size = fw.size_by_key.get(uc.size or "", {})
        ws.append([
            item.computed_rank, uc.ref, uc.title, fw.status_labels.get(uc.status, uc.status), stage["title"] if stage else "",
            uc.domain, uc.use_case_lead, uc.sponsor, uc.story, uc.gap_statement, sol.get(uc.solution_type or "", ""),
            uc.scope_level, uc.effort_level, size.get("label"), size.get("weeks"),
            *[(uc.scores or {}).get(k) for k in fw.lens_keys],
            item.score, uc.agreed_rank, "; ".join(item.flags), uc.qualification_outcome,
            uc.updated_at.strftime("%Y-%m-%d %H:%M") if uc.updated_at else "",
        ])
    header_fill = PatternFill("solid", fgColor=fw.raw["brand"]["colours"]["ink"].lstrip("#"))
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    widths = {"Title": 34, "Story": 60, "Gap statement": 60, "Flags": 34, "Use case lead": 20, "Sponsor": 20, "Status": 18, "Stage": 22, "Domain": 22}
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(h, 12)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
